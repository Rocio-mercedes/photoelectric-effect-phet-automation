#!/usr/bin/env python3
"""
Automatiza los pasos 1-4 de la práctica "Simulaciones numéricas de Interacción de
la Radiación con la Materia: Efecto Fotoeléctrico".

Abre la simulación PhET "Efecto Fotoeléctrico", elige el material (Sodio, Zinc o
Platino), pone la intensidad al 100 % y la fuente en 0 V, busca la longitud de
onda crítica λC con dos criterios —que el amperímetro deje de marcar y que dejen
de verse electrones en el tubo— y después baja la longitud de onda de a 10 nm
veinte veces, midiendo en cada paso el potencial de frenado Vf. Guarda todo en
datos/datos_crudos.xlsx, con una hoja por material y otra con los puntos de cada
recta corriente-voltaje, y al terminar llama a analizar.py, que hace la regresión
de Vf en función de 1/λ y saca φ0.

Cómo se manejan los controles:
    La simulación es un programa Java que corre dentro del navegador y dibuja
    todo como una imagen, así que no hay nada que leer ni que apretar en la
    página. El script usa los controles de verdad —el desplegable "Objetivo" y
    las tres cajas de texto (intensidad, longitud de onda y voltaje)— con clics
    y tecleo, igual que una persona. La caja de la longitud de onda se mueve
    junto con la perilla del espectro, así que su posición se calcula a partir
    de la longitud de onda que haya puesta.
    Para leer el estado se consulta el modelo Java por dentro (ver GANCHO_JS):
    longitud de onda, intensidad, voltaje, función trabajo del material y la
    corriente que muestra el amperímetro. Cada ajuste se verifica contra ese
    estado y se reintenta si no entró, que es lo único que hace fiable el
    tecleo sobre una interfaz dibujada en un lienzo.

Cómo se mide el potencial de frenado:
    Método A (el que pide la práctica): se busca por bisección el |V| más chico
    con el que el amperímetro marca 0, entre 0 y 8 V, en 13 pasos.
    Método B (control): se miden varios puntos de la recta corriente-voltaje y
    se extrapola a corriente cero. Sirve para separar el Vf real del error que
    mete la simulación, que redondea a cero las corrientes muy chicas.

Uso:
    python3 medir_fotoelectrico.py                  # Na, Zn y Pt (~45 min)
    python3 medir_fotoelectrico.py --prueba         # 3 longitudes de onda por material
    python3 medir_fotoelectrico.py --material Na
    python3 medir_fotoelectrico.py --visible        # mostrar la ventana del navegador

Requisitos (una sola vez):
    pip install -r requirements.txt
    (usa el Google Chrome instalado; si no lo hay: python3 -m playwright install chromium)
"""

import argparse
import io
import shutil
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

import matplotlib.image as mpimg
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from playwright.sync_api import Error as ErrorNavegador, sync_playwright

URL = ("https://phet.colorado.edu/sims/cheerpj/photoelectric/latest/"
       "photoelectric.html?simulation=photoelectric&locale=es")

CARPETA = Path(__file__).resolve().parent

# Límites de los controles de la simulación, comprobados uno por uno.
LAMBDA_MIN, LAMBDA_MAX = 100, 850      # nm
V_MAX = 8.0                            # V, la batería no da más
PASO_NM = 10                           # paso que pide la práctica
PASOS = 20                             # cantidad de reducciones que pide la práctica
ITER_BISECCION = 13                    # 8 V / 2^13 ≈ 0.001 V de precisión
PUNTOS_CURVA = 8                       # puntos de la recta I(V) para el método B
PAUSA = 0.15                           # s entre que se escribe algo y se vuelve a leer

# Planilla donde quedan los datos crudos, una hoja por material y otra con los
# puntos de cada recta corriente-voltaje. Los encabezados tienen que coincidir
# con los que espera analizar.py.
ARCHIVO_DATOS = "datos_crudos.xlsx"
COLUMNAS = ["Paso", "λ (nm)", "Vf (V)", "φ₀ simulación (eV)", "λC amperímetro (nm)",
            "λC electrones (nm)", "φ₀ literatura (eV)"]
COLUMNAS_CURVAS = ["λ (nm)", "V (V)", "Corriente"]

# Posiciones (en píxeles) de los controles dentro de la simulación, que siempre
# se dibuja en un lienzo de 1024x768 anclado arriba a la izquierda.
DESPLEGABLE = (888, 174)
ITEM_X, ITEM_Y0, ITEM_DY = 885, 195, 16        # lista desplegada de "Objetivo"
CAJA_INTENSIDAD = (660, 702, 150)              # x inicial, x final, y
CAJA_VOLTAJE = (328, 382, 598)
# La caja que muestra la longitud de onda viaja con la perilla del espectro.
# Su borde izquierdo es lineal en λ: medido sobre capturas, va de x = 402 a
# x = 674 mientras λ va de 100 a 850 nm, y la caja mide 57 px de ancho.
Y_LAMBDA = 183
CAJA_LAMBDA_X0, CAJA_LAMBDA_PENDIENTE, CAJA_LAMBDA_ANCHO = 365.4, 0.3631, 57
PERILLA_INTENSIDAD = (508, 700, 151)           # respaldo: arrastrar el deslizador

# Interior del tubo, entre las dos placas: ahí se ven los electrones arrancados,
# bolitas azul oscuro sobre el fondo blanco y sobre el tinte pastel del haz.
TUBO = {"x": 125, "y": 250, "width": 385, "height": 210}
ELECTRONES_MINIMO = 20                 # píxeles azules: una bolita son ~40
ESPERA_VACIAR = 45                     # s como mucho esperando a que se vacíe el tubo
ESPERA_LLENAR = 4                      # s hasta que se ven los nuevos

MATERIALES = {
    "Na": {"nombre": "Sodio", "item": 0, "phi_lit": 2.46},
    "Zn": {"nombre": "Zinc", "item": 1, "phi_lit": 4.31},
    "Pt": {"nombre": "Platino", "item": 3, "phi_lit": 6.35},
}

# Se inyecta antes de que cargue la simulación. Espía el campo interno que
# alimenta el recuadro "Corriente" y, de paso, se queda con el objeto del modelo,
# desde el cual se llega al resto del estado sin tocar nada.
GANCHO_JS = """
(() => {
  window.__modelo = null;
  window.__corriente = 0;
  function espiar(campo, alCambiar) {
    const oculto = '__' + campo;
    Object.defineProperty(Object.prototype, campo, {
      configurable: true,
      get() { return this[oculto] === undefined ? 0 : this[oculto]; },
      set(v) {
        Object.defineProperty(this, oculto,
          {value: v, writable: true, configurable: true, enumerable: false});
        alCambiar(this, v);
      }
    });
  }
  espiar('wavelength61', (modelo, v) => { window.__modelo = modelo; });
  espiar('current59', (modelo, v) => { window.__corriente = v; });
})();
"""

# Lee el estado del modelo Java: longitud de onda del haz, intensidad como
# fracción del máximo, voltaje entre placas, función trabajo del material
# elegido y la corriente que muestra el amperímetro.
ESTADO_JS = """
(() => {
  const m = window.__modelo;
  if (!m) return null;
  const haz = m.beam50;
  return {
    lambda: haz.wavelength6,
    intensidad: haz.maxPhotonsPerSecond10
      ? haz.photonsPerSecond9 / haz.maxPhotonsPerSecond10 * 100 : 0,
    voltaje: -(m.rightHandPlate49.potential5 - m.target45.potential5),
    phi: m.target45.targetMaterial13.workFunction6,
    corriente: window.__corriente,
  };
})()
"""


class Simulacion:
    """La simulación abierta en el navegador, con un control por cada perilla."""

    def __init__(self, pagina):
        self.pagina = pagina

    # --- lectura ------------------------------------------------------------

    def estado(self):
        return self.pagina.evaluate(ESTADO_JS)

    def corriente(self):
        """Corriente que muestra el amperímetro, después de un par de pasos de reloj."""
        time.sleep(PAUSA)
        anterior = self.pagina.evaluate("window.__corriente")
        limite = time.time() + 1.5
        while time.time() < limite:
            time.sleep(PAUSA)
            actual = self.pagina.evaluate("window.__corriente")
            if actual == anterior:
                return actual
            anterior = actual
        return anterior

    def hay_electrones(self):
        """Si se ven electrones viajando dentro del tubo.

        Es el otro criterio para decir que hay efecto fotoeléctrico, y no es el
        mismo que el del amperímetro: la simulación redondea a cero las
        corrientes chicas, pero los electrones se siguen viendo salir.
        """
        captura = self.pagina.screenshot(clip=TUBO)
        imagen = mpimg.imread(io.BytesIO(captura), format="png")[:, :, :3]
        rojo, verde, azul = imagen[:, :, 0], imagen[:, :, 1], imagen[:, :, 2]
        azul_oscuro = (rojo < 0.5) & (verde < 0.5) & (azul > 0.5)
        return int(np.count_nonzero(azul_oscuro)) >= ELECTRONES_MINIMO

    # --- escritura ----------------------------------------------------------

    def _tipear(self, texto):
        self.pagina.keyboard.type(texto, delay=30)
        time.sleep(PAUSA)
        self.pagina.keyboard.press("Enter")

    def _seleccionar_y_tipear(self, caja, texto):
        """Reemplaza el contenido de una caja de texto, como se hace a mano.

        Primero un clic suelto, que le da el foco: si la caja no lo tiene, se
        come el primer gesto y la selección no llega a hacerse. Después un
        arrastre de punta a punta, que selecciona todo lo que hay escrito, y
        recién ahí el número nuevo.
        """
        x0, x1, y = caja
        self.pagina.mouse.click((x0 + x1) // 2, y)
        time.sleep(PAUSA)
        self.pagina.mouse.move(x0, y)
        self.pagina.mouse.down()
        self.pagina.mouse.move(x1, y, steps=4)
        self.pagina.mouse.up()
        time.sleep(PAUSA)
        self._tipear(texto)

    def _caja_lambda(self):
        """Dónde quedó la caja de λ, que viaja con la perilla del espectro."""
        x = round(CAJA_LAMBDA_X0 + CAJA_LAMBDA_PENDIENTE * self.estado()["lambda"])
        return x + 5, x + CAJA_LAMBDA_ANCHO - 5, Y_LAMBDA

    def _insistir(self, escribir, leer, objetivo, tolerancia, que,
                  intentos=5, sacudir=None, espera=1.5):
        """Escribe un valor y lo verifica contra el modelo; reintenta si no entró.

        `sacudir` se usa cuando la caja ya muestra el número pedido y por eso no
        dispara ningún cambio: escribe otro valor antes de reintentar.
        Entre lectura y lectura hay una pausa holgada a propósito: la simulación
        y el script comparten el único hilo de JavaScript de la página, así que
        preguntarle el estado muy seguido le saca el tiempo que necesita para
        atender el teclado.
        """
        for intento in range(intentos):
            if intento and sacudir:
                sacudir()
                time.sleep(PAUSA)
            escribir()
            limite = time.time() + espera
            while time.time() < limite:
                time.sleep(PAUSA)
                if abs(leer() - objetivo) <= tolerancia:
                    return
        raise RuntimeError(f"No se pudo poner {que} en {objetivo} "
                           f"(quedó en {leer():.4g}). ¿Cambió la interfaz?")

    def longitud_de_onda(self, nm):
        self._insistir(lambda: self._seleccionar_y_tipear(self._caja_lambda(), str(nm)),
                       lambda: self.estado()["lambda"], nm, 0.5,
                       "la longitud de onda")

    def voltaje(self, v):
        """Pone la fuente en v voltios (v ≤ 0 frena a los electrones).

        Pidiéndole 0 V la simulación queda en 0.004 V, que para la medición es
        lo mismo: la corriente sólo depende del voltaje cuando es negativo.
        """
        v = round(v, 3)
        tolerancia = 6e-3 if v == 0 else 5e-4
        otro = v + 1.0 if v + 1.0 <= V_MAX else v - 1.0
        self._insistir(lambda: self._seleccionar_y_tipear(CAJA_VOLTAJE, f"{v:.3f}"),
                       lambda: self.estado()["voltaje"], v, tolerancia, "el voltaje",
                       sacudir=lambda: self._seleccionar_y_tipear(CAJA_VOLTAJE, f"{otro:.3f}"))

    def intensidad_al_maximo(self):
        """Pone la fuente al 100 %: primero por la caja de texto, si no arrastrando."""
        try:
            self._insistir(lambda: self._seleccionar_y_tipear(CAJA_INTENSIDAD, "100"),
                           lambda: self.estado()["intensidad"], 100.0, 0.5,
                           "la intensidad", intentos=3, espera=3.0)
        except RuntimeError:
            x0, x1, y = PERILLA_INTENSIDAD
            def arrastrar():
                self.pagina.mouse.move(x0, y)
                self.pagina.mouse.down()
                self.pagina.mouse.move(x1, y, steps=10)
                self.pagina.mouse.up()
            self._insistir(arrastrar, lambda: self.estado()["intensidad"], 100.0, 0.5,
                           "la intensidad")

    def material(self, clave):
        info = MATERIALES[clave]
        for _ in range(4):
            self.pagina.mouse.click(*DESPLEGABLE)
            time.sleep(0.7)
            self.pagina.mouse.click(ITEM_X, ITEM_Y0 + info["item"] * ITEM_DY)
            time.sleep(0.9)
            phi = self.estado()["phi"]
            if phi > 0:
                return phi
        raise RuntimeError(f"No se pudo elegir {info['nombre']} en el desplegable.")


def abrir(pagina):
    print("Abriendo la simulación (tarda unos 20-40 s en cargar)...")
    pagina.goto(URL, timeout=180_000)
    limite = time.time() + 180
    while time.time() < limite:
        if pagina.evaluate("!!window.__modelo"):
            break
        time.sleep(1)
    else:
        sys.exit("La simulación no terminó de cargar. Probá de nuevo.")
    time.sleep(6)  # margen para que termine de dibujarse la interfaz
    return Simulacion(pagina)


def longitud_critica(sim):
    """Paso 3: máxima λ a la que todavía hay efecto fotoeléctrico, por bisección.

    La corriente crece cuando λ baja, así que alcanza con buscar el borde entre
    "hay corriente" y "no hay corriente" dentro del rango del control.
    """
    sim.voltaje(0)
    sim.longitud_de_onda(LAMBDA_MAX)
    if sim.corriente() > 0:
        return LAMBDA_MAX
    sim.longitud_de_onda(LAMBDA_MIN)
    if sim.corriente() == 0:
        raise RuntimeError("No hay efecto fotoeléctrico ni a 100 nm.")
    con, sin = LAMBDA_MIN, LAMBDA_MAX
    while sin - con > 1:
        medio = (con + sin) // 2
        sim.longitud_de_onda(medio)
        if sim.corriente() > 0:
            con = medio
        else:
            sin = medio
    return con


def vaciar(sim):
    """Espera a ver el tubo vacío. No alcanza con esperar un rato fijo: si queda
    algún electrón de la prueba anterior se cuenta como si fuera nuevo."""
    limite = time.time() + ESPERA_VACIAR
    while sim.hay_electrones():
        if time.time() > limite:
            raise RuntimeError("El tubo no termina de vaciarse. ¿Se frenó la simulación?")
        time.sleep(1)


def longitud_critica_visible(sim):
    """Paso 3, otra vez: la máxima λ a la que todavía salen electrones.

    Mismo bisección, pero mirando el tubo en vez del amperímetro. Antes de cada
    prueba se apaga la emisión poniendo λ = 850 nm y se espera a ver el tubo
    vacío: si uno se confía en esperar un rato fijo, los electrones que quedaron
    viajando de la prueba anterior se cuentan como si fueran nuevos y la
    bisección se va a cualquier lado.
    """
    # Medir los potenciales de frenado deja electrones casi parados contra la
    # placa, que ya no se van solos. Un golpe de +8 V los barre hacia la otra.
    sim.longitud_de_onda(LAMBDA_MAX)
    sim.voltaje(V_MAX)
    vaciar(sim)
    sim.voltaje(0)

    def salen(nm):
        sim.longitud_de_onda(LAMBDA_MAX)        # a 850 nm no emite ningún material
        vaciar(sim)
        sim.longitud_de_onda(nm)
        time.sleep(ESPERA_LLENAR)
        return sim.hay_electrones()

    if not salen(LAMBDA_MIN):
        raise RuntimeError("No salen electrones ni a 100 nm.")
    con, sin = LAMBDA_MIN, LAMBDA_MAX
    while sin - con > 1:
        medio = (con + sin) // 2
        if salen(medio):
            con = medio
        else:
            sin = medio
    return con


def potencial_de_frenado(sim):
    """Paso 4: |V| más chico que lleva el amperímetro a 0, por bisección."""
    sim.voltaje(0)
    if sim.corriente() == 0:
        return 0.0
    pasa, corta = 0.0, V_MAX
    for _ in range(ITER_BISECCION):
        medio = (pasa + corta) / 2
        sim.voltaje(-medio)
        if sim.corriente() > 0:
            pasa = medio
        else:
            corta = medio
    return corta


def curva_corriente_voltaje(sim, vf):
    """Puntos de la recta I(V) dentro del tramo de frenado (método B).

    Se reparten entre 0 y el potencial de frenado sin incluir los extremos: en
    0 V la simulación deja 4 mV de resto y en Vf la corriente ya es cero.
    """
    puntos = []
    for k in range(1, PUNTOS_CURVA + 1):
        v = -vf * k / (PUNTOS_CURVA + 1)
        sim.voltaje(float(v))
        estado = sim.estado()
        puntos.append((round(estado["voltaje"], 4), sim.corriente()))
    return puntos


def medir_material(sim, clave, pasos):
    info = MATERIALES[clave]
    print(f"\n=== {info['nombre']} ===")
    phi = sim.material(clave)
    sim.intensidad_al_maximo()
    sim.voltaje(0)
    print(f"  función trabajo que usa la simulación: {phi:.2f} eV")

    lambda_c = longitud_critica(sim)
    print(f"  longitud de onda crítica (amperímetro en cero): {lambda_c} nm")
    lambda_c_vis = longitud_critica_visible(sim)
    print(f"  longitud de onda crítica (dejan de salir electrones): {lambda_c_vis} nm")

    filas, curvas = [], []
    for k in range(1, pasos + 1):
        nm = lambda_c - PASO_NM * k
        if nm < LAMBDA_MIN:
            print(f"  (la simulación no baja de {LAMBDA_MIN} nm: "
                  f"quedan {k - 1} pasos de los {pasos} pedidos)")
            break
        sim.longitud_de_onda(nm)
        vf = potencial_de_frenado(sim)
        if vf >= V_MAX - 1e-3:
            print(f"  (a {nm} nm haría falta más de {V_MAX:g} V: se corta acá)")
            break
        if vf <= 0:
            print(f"  (a {nm} nm el amperímetro ya marca cero sin frenar: se corta acá)")
            break
        puntos = curva_corriente_voltaje(sim, vf)
        curvas += [(nm, v, i) for v, i in puntos]
        filas.append((k, nm, vf))
        print(f"  {k:2d}/{pasos}   λ = {nm:3d} nm   1/λ = {1 / nm:.6f} nm⁻¹"
              f"   Vf = {vf:.4f} V")
    return {"phi_sim": phi, "lambda_c": lambda_c, "lambda_c_vis": lambda_c_vis,
            "filas": filas, "curvas": curvas}


def respaldar(carpeta):
    """Guarda una copia de la planilla anterior antes de tocarla.

    Es una copia y no un renombre a propósito: midiendo un solo material se
    reescriben sus dos hojas y las de los otros tienen que quedar donde están.
    """
    archivo = carpeta / ARCHIVO_DATOS
    if archivo.exists():
        sello = datetime.fromtimestamp(archivo.stat().st_mtime).strftime("%Y%m%d-%H%M%S")
        shutil.copy2(archivo, carpeta / f"{archivo.stem}_anterior_{sello}.xlsx")


def escribir_hoja(wb, nombre, encabezado, filas):
    if nombre in wb.sheetnames:
        del wb[nombre]
    ws = wb.create_sheet(nombre)
    ws.append(encabezado)
    for celda in ws[1]:
        celda.font = Font(bold=True)
        celda.fill = PatternFill("solid", fgColor="DDE7F3")
        celda.alignment = Alignment(horizontal="center", wrap_text=True)
    for fila in filas:
        ws.append(list(fila))
    for columna, titulo in zip(ws.columns, encabezado):
        ws.column_dimensions[columna[0].column_letter].width = max(10, len(titulo) + 2)
    ws.freeze_panes = "A2"


def ordenar_hojas(wb):
    """Deja las hojas en el orden de los materiales y no en el que se midieron."""
    deseado = [nombre for c in MATERIALES for nombre in (c, f"curvas {c}")
               if nombre in wb.sheetnames]
    for destino, nombre in enumerate(deseado):
        wb.move_sheet(nombre, offset=destino - wb.sheetnames.index(nombre))


def guardar(carpeta, clave, r):
    """Agrega las dos hojas de este material a la planilla de datos crudos."""
    carpeta.mkdir(parents=True, exist_ok=True)
    archivo = carpeta / ARCHIVO_DATOS
    wb = load_workbook(archivo) if archivo.exists() else Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    info = MATERIALES[clave]
    escribir_hoja(wb, clave, COLUMNAS,
                  [(k, nm, round(vf, 4), r["phi_sim"], r["lambda_c"], r["lambda_c_vis"],
                    info["phi_lit"]) for k, nm, vf in r["filas"]])
    escribir_hoja(wb, f"curvas {clave}", COLUMNAS_CURVAS,
                  [(nm, v, round(i, 6)) for nm, v, i in r["curvas"]])
    ordenar_hojas(wb)
    wb.save(archivo)
    print(f"Guardado: {archivo.name}, hojas «{clave}» y «curvas {clave}»")


def main():
    ap = argparse.ArgumentParser(
        description="Mide el potencial de frenado en la simulación PhET del efecto fotoeléctrico.")
    ap.add_argument("--material", choices=list(MATERIALES) + ["todos"], default="todos")
    ap.add_argument("-n", type=int, help=f"reducciones de λ por material (default {PASOS})")
    ap.add_argument("--prueba", action="store_true",
                    help="corrida corta en la carpeta prueba/, 3 longitudes de onda")
    ap.add_argument("--visible", action="store_true", help="mostrar la ventana del navegador")
    args = ap.parse_args()
    sys.stdout.reconfigure(line_buffering=True)  # mostrar el progreso en el momento

    pasos = args.n or (3 if args.prueba else PASOS)
    salida = CARPETA / "prueba" if args.prueba else CARPETA
    claves = list(MATERIALES) if args.material == "todos" else [args.material]

    with sync_playwright() as p:
        try:
            navegador = p.chromium.launch(channel="chrome", headless=not args.visible)
        except Exception:
            navegador = p.chromium.launch(headless=not args.visible)
        pagina = navegador.new_page(viewport={"width": 1024, "height": 768})
        pagina.add_init_script(GANCHO_JS)
        sim = abrir(pagina)
        respaldar(salida / "datos")
        for clave in claves:
            try:
                r = medir_material(sim, clave, pasos)
            except ErrorNavegador:
                sys.exit("\n¡Se cerró el navegador! Volvé a correr el script.")
            if r["filas"]:
                guardar(salida / "datos", clave, r)
        navegador.close()

    print("\nAnalizando...")
    subprocess.run([sys.executable, str(CARPETA / "analizar.py"), str(salida)], check=False)


if __name__ == "__main__":
    main()
