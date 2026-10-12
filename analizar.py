#!/usr/bin/env python3
"""
Pasos 5-6 de la práctica: grafica Vf en función de 1/λ, hace la regresión lineal
y de ahí saca la función trabajo φ0 y la constante de Planck h, con la
diferencia relativa porcentual contra los valores de literatura.

    e·Vf = h·c/λ − φ0      →      Vf = (hc/e)·(1/λ) − φ0/e

La pendiente de Vf contra 1/λ vale hc/e y la ordenada al origen vale −φ0/e.

Lee datos/datos_crudos.xlsx: una hoja por material con la longitud de onda y el
potencial de frenado, y una hoja "curvas <material>" con los puntos de la recta
corriente-voltaje. Se pueden completar a mano si hace falta, respetando los
encabezados.

Dos maneras de sacar Vf de la simulación:
    Método A (el que pide la práctica): el |V| más chico con el que el
        amperímetro marca cero.
    Método B (control): se ajusta la recta I(V) y se extrapola a I = 0. Evita
        el error del método A, porque la simulación redondea a cero las
        corrientes por debajo de 0.015 y entonces corta un poco antes de tiempo.

Genera:
    efecto_fotoelectrico.xlsx   una hoja por material (tabla, gráfico) y una
                                hoja "Resumen" con las regresiones
    graficos/*.png              Vf vs 1/λ por material, los tres juntos y las
                                rectas corriente-voltaje

Uso:
    python3 analizar.py            # usa la carpeta del script
    python3 analizar.py prueba     # usa prueba/datos y escribe ahí los resultados
"""

import sys
from collections import defaultdict
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill
from scipy import stats

CARPETA = Path(__file__).resolve().parent
SALIDA = CARPETA / sys.argv[1] if len(sys.argv) > 1 else CARPETA
DATOS = SALIDA / "datos"
GRAFICOS = SALIDA / "graficos"
EXCEL = SALIDA / "efecto_fotoelectrico.xlsx"

# Planilla de datos crudos y nombres de sus columnas, tal como los escribe
# medir_fotoelectrico.py. Si se editan a mano, tienen que seguir diciendo esto.
ARCHIVO_DATOS = DATOS / "datos_crudos.xlsx"
COL_LAMBDA, COL_VF = "λ (nm)", "Vf (V)"
COL_PHI_SIM, COL_PHI_LIT = "φ₀ simulación (eV)", "φ₀ literatura (eV)"
COL_LC, COL_LC_VIS = "λC amperímetro (nm)", "λC electrones (nm)"
COL_V, COL_I = "V (V)", "Corriente"

# Constantes (CODATA 2018). La simulación usa hc = 1239.80 eV·nm.
CARGA = 1.602176634e-19        # C
LUZ = 2.99792458e17            # nm/s
PLANCK = 6.62607015e-34        # J·s
HC = PLANCK * LUZ / CARGA      # eV·nm

# Resolución con la que queda determinado cada potencial de frenado (el paso al
# que llega la bisección). La longitud de onda no aporta incerteza: el control
# la fija en un número entero de nanómetros, no se lee de una escala.
RES_VF = 0.001                 # V

MATERIALES = {
    "Na": {"nombre": "Sodio", "simbolo": "Na", "color": "tab:orange"},
    "Zn": {"nombre": "Zinc", "simbolo": "Zn", "color": "tab:blue"},
    "Pt": {"nombre": "Platino", "simbolo": "Pt", "color": "tab:green"},
}

METODOS = [("A", "Corte del amperímetro"), ("B", "Extrapolación de I(V)")]

NEGRITA = Font(bold=True)
ENCABEZADO = PatternFill("solid", fgColor="DDE7F3")


def numero(valor):
    """Pasa a número lo que venga de la planilla, ya sea número o texto."""
    if isinstance(valor, (int, float)):
        return float(valor)
    if isinstance(valor, str) and valor.strip():
        try:
            return float(valor.replace(",", "."))
        except ValueError:
            return None
    return None


def leer_hoja(wb, nombre):
    """Devuelve una hoja como lista de diccionarios {encabezado: valor}."""
    if nombre not in wb.sheetnames:
        return []
    filas = list(wb[nombre].values)
    if not filas:
        return []
    encabezado = [str(c).strip() if c is not None else "" for c in filas[0]]
    return [dict(zip(encabezado, fila)) for fila in filas[1:]
            if fila and numero(fila[0]) is not None]


def leer_medidas(wb, clave):
    """Devuelve las λ, los Vf y los datos del material que trae el encabezado."""
    filas = leer_hoja(wb, clave)
    if not filas:
        return None
    primera = filas[0]
    return {
        "lambdas": np.array([numero(f[COL_LAMBDA]) for f in filas]),
        "vf_a": np.array([numero(f[COL_VF]) for f in filas]),
        "phi_sim": numero(primera.get(COL_PHI_SIM)),
        "phi_lit": numero(primera.get(COL_PHI_LIT)),
        "lambda_c": numero(primera.get(COL_LC)),
        "lambda_c_vis": numero(primera.get(COL_LC_VIS)),
    }


def leer_curvas(wb, clave):
    """Devuelve {λ: (V, I)} con los puntos de la recta corriente-voltaje."""
    puntos = defaultdict(list)
    for fila in leer_hoja(wb, f"curvas {clave}"):
        puntos[numero(fila[COL_LAMBDA])].append((numero(fila[COL_V]), numero(fila[COL_I])))
    return {nm: (np.array([v for v, _ in p]), np.array([i for _, i in p]))
            for nm, p in puntos.items()}


def extrapolar(voltajes, corrientes):
    """Vf a partir de la recta I(V): I = a + m·V se anula en V = −a/m.

    Se descartan los puntos con corriente nula y los de corriente máxima, que
    son los del tramo saturado (la simulación le pone un techo a la corriente).
    """
    sirven = (corrientes > 0) & (corrientes < corrientes.max() - 1e-12)
    if sirven.sum() < 3:
        sirven = corrientes > 0
    if sirven.sum() < 2:
        return None
    reg = stats.linregress(voltajes[sirven], corrientes[sirven])
    if reg.slope <= 0:
        return None
    return reg.intercept / reg.slope


def regresion(lambdas, vf):
    """Vf = (hc/e)·(1/λ) − φ0/e. Devuelve pendiente, ordenada, φ0 y h.

    Cada parámetro lleva dos incertezas combinadas en cuadratura: la dispersión
    de los puntos alrededor de la recta y la que propaga la resolución con que
    se midió cada Vf. Para el método del corte manda la primera; para el de
    extrapolación, que cae casi exacto sobre la recta, manda la segunda.
    """
    x = 1 / lambdas
    reg = stats.linregress(x, vf)

    n = len(x)
    sxx = np.sum((x - x.mean()) ** 2)
    err_pendiente = np.hypot(reg.stderr, RES_VF / np.sqrt(sxx))
    err_ordenada = np.hypot(reg.intercept_stderr,
                            RES_VF * np.sqrt(1 / n + x.mean() ** 2 / sxx))

    phi = -reg.intercept                       # eV
    planck = reg.slope * CARGA / LUZ           # J·s
    planck_err = err_pendiente * CARGA / LUZ
    return {
        "n": n, "x": x, "y": vf,
        "pendiente": reg.slope, "pendiente_err": err_pendiente,
        "ordenada": reg.intercept, "ordenada_err": err_ordenada,
        "r2": reg.rvalue ** 2,
        "phi": phi, "phi_err": err_ordenada,
        "h": planck, "h_err": planck_err,
        "dif_h": (planck - PLANCK) / PLANCK * 100,
        # Cuántas veces la incerteza entra en la diferencia con el valor de
        # literatura: si da mucho más que 2, la diferencia no es casualidad.
        "desvio_h": abs(planck - PLANCK) / planck_err,
    }


def diferencia(valor, referencia):
    return (valor - referencia) / referencia * 100


def analizar(wb, clave):
    r = leer_medidas(wb, clave)
    if r is None:
        return None
    if len(r["lambdas"]) < 3:
        print(f"La hoja «{clave}» tiene menos de 3 longitudes de onda, se omite.")
        return None
    lambdas, vf_a = r["lambdas"], r.pop("vf_a")
    r["curvas"] = leer_curvas(wb, clave)
    r["lambda_c_teorica"] = HC / r["phi_sim"] if r["phi_sim"] else None
    r["vf"] = {"A": vf_a}
    r["ajuste"] = {"A": regresion(lambdas, vf_a)}

    vf_b = [extrapolar(*r["curvas"][nm]) if nm in r["curvas"] else None for nm in lambdas]
    if all(v is not None for v in vf_b):
        r["vf"]["B"] = np.array(vf_b, dtype=float)
        r["ajuste"]["B"] = regresion(lambdas, r["vf"]["B"])
    return r


# --- gráficos ---------------------------------------------------------------

def graficar_material(clave, r):
    GRAFICOS.mkdir(parents=True, exist_ok=True)
    info = MATERIALES[clave]
    x = 1 / r["lambdas"]
    recta = np.linspace(0, x.max() * 1.05, 200)

    fig, ax = plt.subplots(figsize=(7, 4.8))
    estilos = {"A": ("o", "tab:blue", "-"), "B": ("s", "tab:red", "--")}
    for metodo, nombre in METODOS:
        if metodo not in r["ajuste"]:
            continue
        a = r["ajuste"][metodo]
        marca, color, linea = estilos[metodo]
        ax.plot(x, r["vf"][metodo], marca, ms=5, color=color, label=f"Vf — {nombre}")
        ax.plot(recta, a["pendiente"] * recta + a["ordenada"], linea, color=color, lw=1.2,
                label=f"   φ₀ = {a['phi']:.3f} eV,  R² = {a['r2']:.5f}")
    ax.axhline(0, color="0.6", lw=0.8)
    ax.set(xlabel="1/λ (nm⁻¹)", ylabel="Potencial de frenado Vf (V)",
           title=f"Vf en función de 1/λ — {info['nombre']}")
    ax.set_xlim(0, None)
    ax.legend(fontsize=8)
    ax.grid(alpha=0.3)
    fig.tight_layout()
    fig.savefig(GRAFICOS / f"Vf_vs_inv_lambda_{clave}.png", dpi=150)
    plt.close(fig)


def graficar_curvas(clave, r):
    """Las rectas corriente-voltaje que se usaron para el método B."""
    if not r["curvas"]:
        return
    info = MATERIALES[clave]
    fig, ax = plt.subplots(figsize=(7, 4.8))
    nms = sorted(r["curvas"])
    colores = plt.cm.viridis(np.linspace(0, 0.9, len(nms)))
    for nm, color in zip(nms, colores):
        v, i = r["curvas"][nm]
        orden = np.argsort(v)
        ax.plot(v[orden], i[orden], "o-", ms=3, lw=1, color=color, label=f"{nm:.0f} nm")
    ax.set(xlabel="V (V)", ylabel="Corriente (unidades de la simulación)",
           title=f"Corriente en función del voltaje — {info['nombre']}")
    ax.legend(fontsize=7, ncol=2, title="λ")
    ax.grid(alpha=0.3)
    fig.tight_layout()
    fig.savefig(GRAFICOS / f"curvas_IV_{clave}.png", dpi=150)
    plt.close(fig)


def graficar_comparacion(resultados):
    """Los tres materiales juntos: rectas paralelas, cada una con su φ0.

    La pendiente común es hc/e y no depende del material; lo único que cambia
    es dónde corta cada recta, que es justamente −φ0/e.
    """
    fig, ax = plt.subplots(figsize=(7.5, 5))
    x_max = max((1 / r["lambdas"]).max() for r in resultados.values()) * 1.05
    for clave, r in resultados.items():
        info = MATERIALES[clave]
        a = r["ajuste"]["A"]
        x = 1 / r["lambdas"]
        ax.plot(x, r["vf"]["A"], "o", ms=5, color=info["color"],
                label=f"{info['nombre']} — φ₀ = {a['phi']:.3f} eV "
                      f"(literatura {r['phi_lit']:g} eV)")
        # Cada recta se dibuja sólo sobre su propio rango, desde donde corta
        # el eje (ahí λ = λC) hasta la λ más corta que se pudo medir.
        corte = -a["ordenada"] / a["pendiente"]
        recta = np.linspace(corte, x.max(), 100)
        ax.plot(recta, a["pendiente"] * recta + a["ordenada"], color=info["color"], lw=1.2)
        ax.plot([corte], [0], "v", ms=7, color=info["color"])
    ax.axhline(0, color="0.5", lw=0.9)
    ax.set(xlabel="1/λ (nm⁻¹)", ylabel="Potencial de frenado Vf (V)",
           title="Vf en función de 1/λ para los tres materiales\n"
                 "(misma pendiente hc/e; la ordenada al origen es −φ₀/e)")
    ax.set_xlim(0, x_max)
    ax.set_ylim(bottom=0)
    ax.legend(fontsize=8)
    ax.grid(alpha=0.3)
    fig.tight_layout()
    fig.savefig(GRAFICOS / "comparacion_materiales.png", dpi=150)
    plt.close(fig)


# --- planilla ---------------------------------------------------------------

def encabezado(ws, fila, columna, textos):
    for i, texto in enumerate(textos):
        celda = ws.cell(fila, columna + i, texto)
        celda.font = NEGRITA
        celda.fill = ENCABEZADO
        celda.alignment = Alignment(horizontal="center", wrap_text=True)


def hoja_material(wb, clave, r):
    info = MATERIALES[clave]
    ws = wb.create_sheet(clave)
    ws["A1"] = f"{info['nombre']} — potencial de frenado en función de la longitud de onda"
    ws["A1"].font = Font(bold=True, size=13)
    visible = (f"    λC con el criterio de los electrones: {r['lambda_c_vis']:.0f} nm"
               if r["lambda_c_vis"] else "")
    ws["A2"] = (f"λC con el amperímetro en cero: {r['lambda_c']:.0f} nm{visible}    "
                f"(hc/φ₀ con la φ₀ de la simulación = {r['lambda_c_teorica']:.1f} nm)")
    ws["A3"] = (f"φ₀ de la simulación: {r['phi_sim']:g} eV    "
                f"φ₀ de literatura (Serway, tabla 40.1): {r['phi_lit']:g} eV")

    columnas = ["Paso", "λ (nm)", "1/λ (nm⁻¹)", "Vf — corte (V)"]
    if "B" in r["vf"]:
        columnas.append("Vf — extrapolación (V)")
    encabezado(ws, 5, 1, columnas)
    for i, nm in enumerate(r["lambdas"]):
        fila = 6 + i
        ws.cell(fila, 1, i + 1)
        ws.cell(fila, 2, int(nm))
        ws.cell(fila, 3, round(float(1 / nm), 8))
        ws.cell(fila, 4, round(float(r["vf"]["A"][i]), 4))
        if "B" in r["vf"]:
            ws.cell(fila, 5, round(float(r["vf"]["B"][i]), 4))

    ultima = 5 + len(r["lambdas"])
    for col, ancho in zip("ABCDE", [7, 9, 13, 15, 19]):
        ws.column_dimensions[col].width = ancho

    graf = ScatterChart()
    graf.title = f"Vf vs 1/λ — {info['simbolo']}"
    graf.style = 13
    graf.x_axis.title = "1/λ (nm⁻¹)"
    graf.y_axis.title = "Vf (V)"
    serie = Series(Reference(ws, min_col=4, min_row=6, max_row=ultima),
                   Reference(ws, min_col=3, min_row=6, max_row=ultima), title="Vf")
    serie.marker.symbol = "circle"
    serie.marker.size = 6
    serie.graphicalProperties.line.noFill = True
    graf.series.append(serie)
    graf.legend = None
    graf.width, graf.height = 16, 9
    ws.add_chart(graf, "G5")


def hoja_resumen(wb, resultados):
    ws = wb.create_sheet("Resumen", 0)
    ws["A1"] = "Resumen — Efecto fotoeléctrico (simulación PhET)"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Vf = (hc/e)·(1/λ) − φ₀/e      pendiente = hc/e      ordenada al origen = −φ₀/e"
    ws["A3"] = "Δ (%) = (valor experimental − valor de referencia) / valor de referencia × 100"
    ws["A4"] = ("Errores: dispersión del ajuste combinada con la propagación de δVf = "
                "0,001 V.   σ = veces que la incerteza entra en la diferencia con h de literatura.")

    columnas = ["Material", "Método", "n", "Pendiente (V·nm)", "error", "R²",
                "φ₀ exp (eV)", "error φ₀", "φ₀ simulación (eV)", "Δ vs simulación (%)",
                "φ₀ literatura (eV)", "Δ vs literatura (%)", "h exp (J·s)", "error h",
                "Δ vs h (%)", "desvío de h (σ)"]
    encabezado(ws, 5, 1, columnas)
    fila = 6
    for clave, r in resultados.items():
        for metodo, nombre in METODOS:
            if metodo not in r["ajuste"]:
                continue
            a = r["ajuste"][metodo]
            valores = [MATERIALES[clave]["simbolo"], nombre, a["n"], a["pendiente"],
                       a["pendiente_err"], a["r2"], a["phi"], a["phi_err"],
                       r["phi_sim"], diferencia(a["phi"], r["phi_sim"]),
                       r["phi_lit"], diferencia(a["phi"], r["phi_lit"]),
                       a["h"], a["h_err"], a["dif_h"], a["desvio_h"]]
            for c, v in enumerate(valores, 1):
                ws.cell(fila, c, float(v) if isinstance(v, (np.floating, float)) else v)
            fila += 1

    fila += 1
    ws.cell(fila, 1, "Longitud de onda crítica (paso 3)").font = NEGRITA
    fila += 1
    encabezado(ws, fila, 1, ["Material", "λC amperímetro (nm)", "Δ (%)",
                             "λC electrones (nm)", "Δ (%)", "hc/φ₀ simulación (nm)"])
    for clave, r in resultados.items():
        fila += 1
        ws.cell(fila, 1, MATERIALES[clave]["simbolo"])
        ws.cell(fila, 2, float(r["lambda_c"]))
        ws.cell(fila, 3, float(diferencia(r["lambda_c"], r["lambda_c_teorica"])))
        if r["lambda_c_vis"]:
            ws.cell(fila, 4, float(r["lambda_c_vis"]))
            ws.cell(fila, 5, float(diferencia(r["lambda_c_vis"], r["lambda_c_teorica"])))
        ws.cell(fila, 6, float(r["lambda_c_teorica"]))

    for col, ancho in zip("ABCDEFGHIJKLMNOP",
                          [9, 22, 5, 15, 10, 9, 12, 10, 14, 14, 14, 14, 13, 11, 10, 12]):
        ws.column_dimensions[col].width = ancho
    for filas in ws.iter_rows(min_row=6, max_row=fila):
        for celda in filas:
            if isinstance(celda.value, float):
                celda.number_format = "0.0000" if abs(celda.value) >= 0.01 else "0.000E+00"


def main():
    if not ARCHIVO_DATOS.exists():
        print(f"No está {ARCHIVO_DATOS}. Corré primero medir_fotoelectrico.py")
        return
    wb_datos = load_workbook(ARCHIVO_DATOS, data_only=True)
    resultados = {}
    for clave, info in MATERIALES.items():
        r = analizar(wb_datos, clave)
        if r is None:
            continue
        resultados[clave] = r
        graficar_material(clave, r)
        graficar_curvas(clave, r)

        print(f"\n{info['nombre']}  (n = {r['ajuste']['A']['n']} longitudes de onda)")
        print(f"  λC (amperímetro) = {r['lambda_c']:.0f} nm "
              f"({diferencia(r['lambda_c'], r['lambda_c_teorica']):+.1f} %)", end="")
        if r["lambda_c_vis"]:
            print(f"   λC (electrones) = {r['lambda_c_vis']:.0f} nm "
                  f"({diferencia(r['lambda_c_vis'], r['lambda_c_teorica']):+.1f} %)", end="")
        print(f"   hc/φ₀ = {r['lambda_c_teorica']:.1f} nm")
        for metodo, nombre in METODOS:
            if metodo not in r["ajuste"]:
                continue
            a = r["ajuste"][metodo]
            print(f"  {nombre:24s} φ₀ = {a['phi']:.4f} ± {a['phi_err']:.4f} eV"
                  f"   h = {a['h']:.4e} ± {a['h_err']:.1e} J·s"
                  f" ({a['dif_h']:+.2f} %, {a['desvio_h']:.1f} σ)"
                  f"   R² = {a['r2']:.5f}")
            print(f"  {'':24s}   Δφ₀ = {diferencia(a['phi'], r['phi_sim']):+.2f} % "
                  f"vs simulación ({r['phi_sim']:g} eV),"
                  f" {diferencia(a['phi'], r['phi_lit']):+.2f} % "
                  f"vs literatura ({r['phi_lit']:g} eV)")

    if not resultados:
        print(f"{ARCHIVO_DATOS.name} no tiene datos suficientes de ningún material.")
        return

    graficar_comparacion(resultados)

    wb = Workbook()
    wb.remove(wb.active)
    hoja_resumen(wb, resultados)
    for clave, r in resultados.items():
        hoja_material(wb, clave, r)
    wb.save(EXCEL)
    print(f"\nPlanilla: {EXCEL.name}    Gráficos: {GRAFICOS.name}/")


if __name__ == "__main__":
    main()
